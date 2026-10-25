"""Crossing-house allocation model (v1).

Assigns planned crosses to rooms in a crossing house for a single day.

Assumptions (v1) -- each one is a question for the breeders:
  * One "clone" in a room = one flowering stalk/tassel of that parent, and it
    consumes one unit of the parent's NumberAvailable.
  * Each room holds exactly one male clone (possibly several stalks of it);
    every female stalk in the room is pollinated by that male, so a female is
    only worth placing in a room whose male is a desired partner.
  * Ratio: female stalks in a room <= ratio * male stalks in that room.
  * Capacity: male stalks + female stalks in a room <= capacity.
  * A cross counts as "made" once >= min_female_stalks female stalks are placed
    with its male (default 1).
  * Priority 1 is the most important, 3 the least.
"""

from __future__ import annotations

import argparse
import random
from collections import defaultdict
from dataclasses import dataclass, field

import openpyxl
import pulp

MALE, FEMALE = 1, 2


@dataclass(frozen=True)
class Cross:
    male: str
    female: str
    priority: int

    @property
    def key(self) -> tuple[str, str]:
        return (self.male, self.female)


@dataclass
class Instance:
    crosses: list[Cross]
    available: dict[str, int]  # parent -> flowers available today
    sex: dict[str, int]  # parent -> MALE / FEMALE
    rooms: int = 5
    capacity: int = 8
    ratio: int = 3  # max female stalks per male stalk
    min_female_stalks: int = 1
    warnings: list[str] = field(default_factory=list)


@dataclass
class Room:
    male: str
    male_stalks: int
    females: dict[str, int]  # female parent -> stalks

    @property
    def size(self) -> int:
        return self.male_stalks + sum(self.females.values())


@dataclass
class Solution:
    status: str
    value: float
    rooms: list[Room]
    made: list[Cross]
    unmade: dict[Cross, str]  # cross -> reason


# --------------------------------------------------------------------------- #
# Input
# --------------------------------------------------------------------------- #

def read_workbook(src) -> tuple[list[tuple], dict[str, int], dict[str, int]]:
    """Read DesiredCrosses + ParentInventory from a path or file object,
    using Excel's cached values.

    The toy workbook fills Priority/NumberAvailable with RANDBETWEEN, so the
    cached values are one frozen draw; use `resample` for other draws.
    """
    wb = openpyxl.load_workbook(src, data_only=True)
    for sheet in ("DesiredCrosses", "ParentInventory"):
        if sheet not in wb.sheetnames:
            raise ValueError(f"workbook has no '{sheet}' sheet")
    raw_crosses = [r[:3] for r in wb["DesiredCrosses"].iter_rows(min_row=2, values_only=True) if r[0]]
    inv = [r[:3] for r in wb["ParentInventory"].iter_rows(min_row=2, values_only=True) if r[0]]
    if any(r[2] is None for r in raw_crosses) or any(r[2] is None for r in inv):
        raise ValueError("workbook has formulas without saved values; open and save it in Excel first")
    sex = {str(p): int(s) for p, s, _ in inv}
    available = {str(p): int(n) for p, _, n in inv}
    return raw_crosses, available, sex


def load_instance(path: str, **params) -> Instance:
    return validate(*read_workbook(path), **params)


def validate(raw_crosses, available, sex, **params) -> Instance:
    """Clean the cross list and record every data problem as a warning."""
    warnings = []
    best: dict[tuple[str, str], Cross] = {}
    for m, f, p in raw_crosses:
        c = Cross(str(m), str(f), int(p))
        if c.key in best:
            keep = min(best[c.key], c, key=lambda x: x.priority)
            warnings.append(
                f"{c.male} × {c.female} is listed twice (priority {best[c.key].priority} and {c.priority}); "
                f"using priority {keep.priority}.")
            best[c.key] = keep
        else:
            best[c.key] = c
    crosses = []
    for c in best.values():
        if c.male not in sex or c.female not in sex:
            warnings.append(f"{c.male} × {c.female} was skipped: a parent is not in the flower inventory (dropped).")
            continue
        if sex[c.male] != MALE or sex[c.female] != FEMALE:
            warnings.append(f"{c.male} × {c.female} was skipped: the male/female roles do not match the inventory (dropped).")
            continue
        if c.priority not in (1, 2, 3):
            warnings.append(f"{c.male} × {c.female} has priority {c.priority}; priorities should be 1, 2 or 3.")
        crosses.append(c)
    used = {p for c in crosses for p in c.key}
    idle = sorted(p for p, n in available.items() if p not in used and n > 0)
    if idle:
        warnings.append(f"These parents are flowering but are not in any planned cross: {', '.join(idle)}.")
    return Instance(crosses, available, sex, warnings=warnings, **params)


def resample(inst: Instance, rng: random.Random) -> Instance:
    """New draw of the toy workbook's RANDBETWEEN columns, same structure."""
    crosses = [Cross(c.male, c.female, rng.randint(1, 3)) for c in inst.crosses]
    available = {p: rng.randint(0, 10) for p in inst.available}
    return Instance(crosses, available, inst.sex, inst.rooms, inst.capacity, inst.ratio,
                    inst.min_female_stalks)


# --------------------------------------------------------------------------- #
# Model
# --------------------------------------------------------------------------- #

def priority_weights(crosses: list[Cross], scheme: str) -> dict[int, int]:
    """'linear': 1->3, 2->2, 3->1 (two priority-2 crosses beat one priority-1).
    'strict': lexicographic -- one cross of a better priority outweighs all
    crosses of worse priorities combined."""
    if scheme == "linear":
        return {1: 3, 2: 2, 3: 1}
    if scheme == "strict":
        count = defaultdict(int)
        for c in crosses:
            count[c.priority] += 1
        w = {3: 1}
        w[2] = count[3] * w[3] + 1
        w[1] = count[2] * w[2] + count[3] * w[3] + 1
        return w
    raise ValueError(f"unknown weight scheme {scheme!r}")


def solve(inst: Instance, weights: str = "linear", fill: bool = True) -> Solution:
    """Two-stage lexicographic MIP.

    Stage 1 maximizes the priority-weighted number of crosses made.
    Stage 2 holds that value and, if `fill`, maximizes female stalks placed
    (more pollinations/seed) minus male stalks used; otherwise it minimizes
    total stalks used so flowers are saved for later days. Ties go to the
    allocation using fewer rooms.
    """
    R = range(inst.rooms)
    K, rho = inst.capacity, inst.ratio
    C = [c for c in inst.crosses if inst.available[c.male] > 0
         and inst.available[c.female] >= inst.min_female_stalks]
    males = sorted({c.male for c in C})
    females = sorted({c.female for c in C})
    w = priority_weights(inst.crosses, weights)

    prob = pulp.LpProblem("crossing_house", pulp.LpMaximize)
    # z[m,r] = 1 if room r is assigned male clone m
    z = {(m, r): pulp.LpVariable(f"z_{m}_{r}", cat="Binary") for m in males for r in R}
    # u[m,r] = male stalks of m in room r
    u = {(m, r): pulp.LpVariable(f"u_{m}_{r}", 0, K, cat="Integer") for m in males for r in R}
    # x[c,r] = female stalks placed in room r for cross c
    x = {(c, r): pulp.LpVariable(f"x_{c.male}_{c.female}_{r}", 0, K, cat="Integer") for c in C for r in R}
    # y[c] = 1 if cross c is made
    y = {c: pulp.LpVariable(f"y_{c.male}_{c.female}", cat="Binary") for c in C}

    for r in R:
        prob += pulp.lpSum(z[m, r] for m in males) <= 1, f"one_male_{r}"
        for m in males:
            prob += u[m, r] <= K * z[m, r], f"u_link_{m}_{r}"
            prob += u[m, r] >= z[m, r], f"u_min_{m}_{r}"
        for c in C:
            prob += x[c, r] <= K * z[c.male, r], f"x_link_{c.male}_{c.female}_{r}"
        females_in_r = pulp.lpSum(x[c, r] for c in C)
        males_in_r = pulp.lpSum(u[m, r] for m in males)
        prob += females_in_r <= rho * males_in_r, f"ratio_{r}"
        prob += females_in_r + males_in_r <= K, f"capacity_{r}"
        if r > 0:  # symmetry breaking: use rooms in order
            prob += (pulp.lpSum(z[m, r] for m in males)
                     <= pulp.lpSum(z[m, r - 1] for m in males)), f"order_{r}"
    for m in males:
        prob += pulp.lpSum(u[m, r] for r in R) <= inst.available[m], f"inv_{m}"
    for f in females:
        prob += pulp.lpSum(x[c, r] for c in C if c.female == f for r in R) <= inst.available[f], f"inv_{f}"
    for c in C:
        prob += inst.min_female_stalks * y[c] <= pulp.lpSum(x[c, r] for r in R), f"made_{c.male}_{c.female}"

    solver = pulp.HiGHS(msg=False)
    primary = pulp.lpSum(w[c.priority] * y[c] for c in C)
    prob.setObjective(primary)
    prob.solve(solver)
    status = pulp.LpStatus[prob.status]
    if status != "Optimal":
        return Solution(status, 0, [], [], {c: "solver failed" for c in inst.crosses})
    best = round(pulp.value(primary))

    prob += primary >= best, "hold_primary"
    stalks_f = pulp.lpSum(x.values())
    stalks_m = pulp.lpSum(u.values())
    # Rooms are worth keeping free for tomorrow; the penalty totals < 1 so it
    # only breaks ties between allocations with the same stalk counts.
    rooms_used = pulp.lpSum(z.values()) / (inst.rooms + 1)
    prob.setObjective((stalks_f - stalks_m if fill else -(stalks_f + stalks_m)) - rooms_used)
    prob.solve(solver)

    rooms = []
    for r in R:
        for m in males:
            if z[m, r].value() > 0.5:
                fem = {c.female: round(x[c, r].value()) for c in C
                       if c.male == m and x[c, r].value() > 0.5}
                rooms.append(Room(m, round(u[m, r].value()), fem))
    made = [c for c in C if y[c].value() > 0.5]
    used = flowers_used(rooms)
    unmade = {c: _reason(inst, c, used, len(rooms)) for c in inst.crosses if c not in made}
    return Solution(pulp.LpStatus[prob.status], best, rooms, made, unmade)


def flowers_used(rooms: list[Room]) -> dict[str, int]:
    used = defaultdict(int)
    for room in rooms:
        used[room.male] += room.male_stalks
        for f, n in room.females.items():
            used[f] += n
    return used


def _reason(inst: Instance, c: Cross, used: dict[str, int], rooms_used: int) -> str:
    """Best-guess explanation of why an unmade cross lost out."""
    out = [p for p in c.key if inst.available[p] == 0]
    if out:
        return f"No flowers available on {' or '.join(out)}"
    if inst.available[c.female] < inst.min_female_stalks:
        return f"{c.female} has fewer than the {inst.min_female_stalks} flowers needed"
    spent = [p for p in c.key if used[p] >= inst.available[p]]
    if spent:
        return f"Competition: all {' and '.join(spent)} flowers went to more valuable crosses"
    if rooms_used == inst.rooms:
        return f"Competition: no free room left for male {c.male}"
    return "Competition: other crosses were worth more"


# --------------------------------------------------------------------------- #
# Output
# --------------------------------------------------------------------------- #

def report(inst: Instance, sol: Solution) -> str:
    lines = [f"Status: {sol.status}   rooms {inst.rooms} x {inst.capacity} clones, "
             f"ratio 1 male : {inst.ratio} female"]
    for wmsg in inst.warnings:
        lines.append(f"  ! {wmsg}")
    lines.append("")
    for i, room in enumerate(sol.rooms, 1):
        fem = ", ".join(f"{f} x{n}" for f, n in room.females.items()) or "(none)"
        lines.append(f"  Room {i}: male {room.male} x{room.male_stalks} | females {fem} "
                     f"| {room.size}/{inst.capacity} used")
    for i in range(len(sol.rooms) + 1, inst.rooms + 1):
        lines.append(f"  Room {i}: empty")
    lines.append("")
    lines.append(f"Crosses made ({len(sol.made)}/{len(inst.crosses)}), weighted value {sol.value}:")
    for c in sorted(sol.made, key=lambda c: (c.priority, c.key)):
        lines.append(f"  P{c.priority} {c.male} x {c.female}")
    if sol.unmade:
        lines.append("Not made:")
        for c, why in sorted(sol.unmade.items(), key=lambda kv: (kv[0].priority, kv[0].key)):
            lines.append(f"  P{c.priority} {c.male} x {c.female}: {why}")
    used = flowers_used(sol.rooms)
    lines.append("Flowers used / available: " + ", ".join(
        f"{p} {used[p]}/{inst.available[p]}" for p in sorted(used)))
    return "\n".join(lines)


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("xlsx", nargs="?", default="toyexample.xlsx")
    ap.add_argument("--rooms", type=int, default=5)
    ap.add_argument("--capacity", type=int, default=8)
    ap.add_argument("--ratio", type=int, default=3)
    ap.add_argument("--min-female-stalks", type=int, default=1)
    ap.add_argument("--weights", choices=["linear", "strict"], default="linear")
    ap.add_argument("--conserve", action="store_true",
                    help="use as few flowers as possible instead of filling rooms")
    a = ap.parse_args()
    inst = load_instance(a.xlsx, rooms=a.rooms, capacity=a.capacity, ratio=a.ratio,
                         min_female_stalks=a.min_female_stalks)
    print(report(inst, solve(inst, a.weights, fill=not a.conserve)))


if __name__ == "__main__":
    main()
